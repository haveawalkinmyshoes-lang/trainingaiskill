"""Read an employee workbook locally; never publish the source or generated roster."""
import argparse,json,re,unicodedata
from pathlib import Path
from openpyxl import load_workbook

parser=argparse.ArgumentParser()
parser.add_argument('workbook',type=Path)
parser.add_argument('--sheet')
parser.add_argument('--column',type=int,default=1)
parser.add_argument('--email-column',type=int)
parser.add_argument('--require-email',action='store_true')
parser.add_argument('--output-dir',type=Path,default=Path(__file__).resolve().parents[1]/'private')
args=parser.parse_args()
if args.workbook.suffix.lower()!='.xlsx':
    parser.error('Нужен файл .xlsx. Сохраните старый .xls в формате .xlsx.')
workbook=load_workbook(args.workbook,read_only=True,data_only=True)
sheet=workbook[args.sheet] if args.sheet else workbook.active
headers={'фио','ф.и.о.','сотрудник','сотрудники','фамилия имя','фамилия имя отчество','фамилия и имя','список сотрудников','полное имя','full name','имя пользователя'}
rows=list(sheet.iter_rows(values_only=True))
email_column=args.email_column
if rows:
    for i,value in enumerate(rows[0],1):
        label=str(value or '').strip().lower()
        if label in headers:args.column=i
        if label in {'email','e-mail','почта','электронная почта'}:email_column=i
if args.require_email and not email_column:raise ValueError('Не найден столбец Email')
people=[];seen={};emails=set();duplicates=0
for row_number,row in enumerate(rows,1):
    if len(row)<args.column:continue
    value=row[args.column-1]
    if value is None or not str(value).strip():continue
    if not isinstance(value,str):raise ValueError(f'Строка {row_number}: ФИО должно быть текстом')
    full_name=re.sub(r'\s+',' ',unicodedata.normalize('NFKC',value)).strip()
    normalized=full_name.lower().replace('ё','е')
    if normalized in headers:continue
    if len(full_name)>200 or len(full_name.split())<2 or any(c in full_name for c in '#\x00\r\n'):
        raise ValueError(f'Строка {row_number}: проверьте формат полного имени')
    email=str(row[email_column-1] or '').strip() if email_column and len(row)>=email_column else ''
    if email and not re.fullmatch(r'[^\s@<>]+@[^\s@<>]+\.[^\s@<>]+',email):raise ValueError(f'Строка {row_number}: некорректный Email')
    if args.require_email and not email:raise ValueError(f'Строка {row_number}: не заполнен Email')
    if normalized in seen:
        if seen[normalized]!=email.lower():raise ValueError(f'Строка {row_number}: одинаковые ФИО с разной почтой; уточните полные имена')
        duplicates+=1;continue
    if email and email.lower() in emails:raise ValueError(f'Строка {row_number}: один адрес используется для нескольких сотрудников')
    seen[normalized]=email.lower();emails.add(email.lower()) if email else None
    people.append({'fullName':full_name,'email':email} if email_column else full_name)
workbook.close()
if not people:raise ValueError('В выбранном столбце не найден список сотрудников')
args.output_dir.mkdir(parents=True,exist_ok=True)
encoded=json.dumps(people,ensure_ascii=False,separators=(',',':'))
(args.output_dir/'employees.json').write_text(encoded+'\n',encoding='utf-8')
(args.output_dir/'employees.env').write_text('EMPLOYEE_ALLOWLIST_JSON='+encoded+'\n',encoding='utf-8')
print(json.dumps({'employees':len(people),'with_email':sum(bool(p.get('email')) for p in people if isinstance(p,dict)),'duplicates_merged':duplicates,'source_sheet':sheet.title,'output_dir':str(args.output_dir)},ensure_ascii=False))
